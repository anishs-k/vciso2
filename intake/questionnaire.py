"""The client onboarding questionnaire: an Excel workbook the client fills in, read back into CTW Admin's form.

Asked the way a client can answer:
  "Company Details"  questions and answers (Category | Question | Client Response | Instructions), like
                     QA/Client_Onboarding_Questionnaire.xlsx: company, contract dates, users needed,
                     platform features, sign-in and scanner tools.
  "Offices"          one office per row: type, address, city, state, country, pincode, and the frameworks
                     that office needs with how often they are reviewed.
  "Contacts"         one contact per row.     "Users"  one user per row.
  "Frameworks List"  the frameworks a client can name.
Every answer cell carries a hidden key (column E on "Company Details", row 1 on the tables), so reading does
not depend on the wording. What CTW decides itself (customer and contract IDs, PO number, contract type
and value, currency, invoicing) is not asked; the licence's framework count follows from the answers.

    write_template(path, catalogue)  -> the empty questionnaire (catalogue: {"options": ..., "frameworks": [...]})
    read_answers(path)               -> {"answers": {key: value}, "warnings": [...]}

Keys: "org.customer_name", "office.2.city", "office.2.frameworks", "contact.1.contact_email", "user.3.role_type"...
Values come back as text, numbers as whole numbers where they are whole, dates as YYYY-MM-DD. Mapping
them onto the form (IDs, offices by number) is the UI's job (ctw-admin/ui/src/lib/questionnaire.js).
"""
import base64
import datetime
import json
import re
import sys

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

QA_SHEET = "Company Details"
OLD_SHEET = "Client Onboarding Form"  # the single-sheet questionnaire before the tables
FRAMEWORK_SHEET = "Frameworks List"
OFFICES, CONTACTS, USERS = 5, 8, 8
TABLES = [("Offices", "office"), ("Contacts", "contact"), ("Users", "user")]  # sheet, record key
MODULES = [
    ("vulnerabilities", "Vulnerability Management"),
    ("alerting", "Alerting"),
    ("mobile_access", "Mobile Access"),
    ("news_breach_watch", "News & Breach Watch"),
    ("onprem_ai_assessment", "On-prem AI Assessment"),
    ("incidents", "Incidents"),
]
YES_NO = ["Yes", "No"]
REVIEW_MONTHS = ["1", "3", "6", "12", "24", "36"]

HEAD_FONT = Font(bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", fgColor="1F3A5F")
SECTION_FILL = PatternFill("solid", fgColor="DCE6F1")
ANSWER_FILL = PatternFill("solid", fgColor="FFF9E6")
WRAP = Alignment(vertical="top", wrap_text=True)


def office_names(n=OFFICES):
    return [f"Office {i}" for i in range(1, n + 1)]


def short_name(name):
    """A framework's name without its explanation: "ISO/IEC 27001:2022 (Information ...)" -> "ISO/IEC 27001:2022"."""
    return name.split(" (")[0].strip()


def questions(catalogue):
    """The "Company Details" sheet, top to bottom: ("section", title) or (key, question, help, choices).
    Choices that are a list become a dropdown; ("free", list) only suggests, other text is let in too."""
    opt = catalogue["options"]
    out = [("section", "1. Company Details"),
           ("org.customer_name", "What is your company's registered (legal) name?", "As on your registration documents", None),
           ("org.category", "Which category does your company fall in?", "Choose one", opt["categories"]),
           ("org.industry", "Which industry are you in?", "e.g., Banking & Financial Services, IT, Healthcare", None),
           ("org.primary_domain", "What is your company's main website domain?", "e.g., acmefinance.com", None),
           ("section", "2. Contract"),
           ("contract.contract_start_date", "From which date should the contract start?", "Format: DD-MMM-YYYY (e.g., 01-Apr-2026)", None),
           ("contract.contract_period_months", "For how many months do you need the contract?", "e.g., 12, 24, 36", None),
           ("section", "3. Users Needed"),
           ("licence.max_users", "How many people will use the platform in total?", "All users together", None),
           ("licence.max_vcisos", "How many of them are vCISOs?", "e.g., 1", None),
           ("licence.max_org_owners", "How many Organisation Owners?", "e.g., 1", None),
           ("licence.max_it_admins", "How many IT Admins?", "At least 1", None),
           ("licence.max_api_integrations", "How many scanner tools do you want to connect?", "0 if none", None),
           ("section", "4. Platform Features")]
    for key, label in MODULES:
        out.append((f"module.{key}", f"Do you need {label}?", "Yes or No (Compliance is always included)", YES_NO))
    out += [("news.monitored_names", "News & Breach Watch - Which names should we watch the news for?", "Only if you chose News & Breach Watch; separate with commas", None),
            ("news.monitored_domains", "News & Breach Watch - Which domains should we watch?", "Only if you chose News & Breach Watch; separate with commas", None),
            ("section", "5. Login & Security"),
            ("security.sso_type", "How will your users sign in?", "LOCAL (user ID and password), or your single sign-on type", opt["sso_types"]),
            ("security.sso_metadata_url", "Single sign-on metadata URL", "Only for single sign-on", None),
            ("security.session_timeout_minutes", "After how many idle minutes should users be signed out?", "e.g., 30", None),
            ("security.mfa_enforced", "Should every user sign in with a one-time code (MFA)?", "Yes or No", YES_NO),
            ("section", "6. Scanner Tools")]
    for tool in opt["tools"]:
        out += [(f"tool.{tool}.used", f"Do you use {tool}?", "Yes or No", YES_NO),
                (f"tool.{tool}.endpoint", f"{tool} - Endpoint URL", "Only if you use it, e.g., https://cloud.tenable.com", None)]
    return out


def tables(catalogue):
    """The one-row-per-record sheets: (sheet, record key, rows, what to ask, [(field, heading, width, choices)])."""
    opt = catalogue["options"]
    return [
        ("Offices", "office", OFFICES,
         "Please list all your offices, one office per row. Mark exactly one office as Primary. For each office, "
         f'name the frameworks it needs, separated by commas (all names are on the "{FRAMEWORK_SHEET}" sheet).',
         [("office_type", "Office Type", 18, opt["office_types"]),
          ("office_address", "Address Line", 40, None),
          ("city", "City", 16, None),
          ("state", "State", 18, None),
          ("country", "Country", 14, None),
          ("pincode", "Pincode/Zip Code", 14, None),
          ("frameworks", "Frameworks Needed (e.g., ISO 27001, DPDPA)", 36, None),
          ("review_months", "Review Frameworks Every (months)", 16, REVIEW_MONTHS)]),
        ("Contacts", "contact", CONTACTS,
         "Who should we call, and in what order? One person per row. Office = the office's number on the Offices sheet.",
         [("contact_type", "Contact For", 16, opt["contact_types"]),
          ("contact_priority", "Call Order (1 = first)", 12, None),
          ("office", "Office", 11, office_names()),
          ("contact_person", "Full Name", 22, None),
          ("designation", "Designation", 22, ("free", opt.get("designations", []))),
          ("department", "Department", 24, ("free", opt.get("departments", []))),
          ("contact_email", "Email", 30, None),
          ("contact_phone", "Phone (with country code)", 20, None)]),
        ("Users", "user", USERS,
         "The first people to sign in to the platform. One person per row; at least one must be an IT Admin. "
         "Office = the office's number on the Offices sheet.",
         [("user_name", "Full Name", 22, None),
          ("login_id", "Login ID (e.g., rahul.sharma)", 22, None),
          ("user_email", "Email", 30, None),
          ("user_mobile", "Mobile (with country code)", 20, None),
          ("designation", "Designation", 22, ("free", opt.get("designations", []))),
          ("department", "Department", 24, ("free", opt.get("departments", []))),
          ("role_type", "Role", 18, opt["role_types"]),
          ("office", "Office", 11, office_names())]),
    ]


class Dropdowns:
    """Dropdown lists, kept on a hidden "Lists" sheet (inline lists are limited to 255 characters)."""

    def __init__(self, wb):
        self.lists = wb.create_sheet("Lists")
        self.lists.sheet_state = "hidden"
        self.cols = {}

    def add(self, ws, cells, choices):
        free = isinstance(choices, tuple)
        values = choices[1] if free else choices
        if not values:
            return
        key = tuple(values)
        if key not in self.cols:
            n = len(self.cols) + 1
            for r, v in enumerate(values, 1):
                self.lists.cell(r, n, v)
            self.cols[key] = self.lists.cell(1, n).column_letter
        col = self.cols[key]
        dv = DataValidation(type="list", formula1=f"=Lists!${col}$1:${col}${len(values)}", allow_blank=True)
        dv.showErrorMessage = not free  # a free list only suggests
        ws.add_data_validation(dv)
        dv.add(cells)


def head(cells):
    for c in cells:
        c.font = HEAD_FONT
        c.fill = HEAD_FILL
        c.alignment = WRAP


def write_template(path, catalogue):
    wb = Workbook()
    ws = wb.active
    ws.title = QA_SHEET
    drops = Dropdowns(wb)

    ws.append(["Category", "Question", "Client Response (Please fill here)", "Instructions / Examples", "Key"])
    head(ws[1])
    section = ""
    for row in questions(catalogue):
        if row[0] == "section":
            section = row[1].split(". ", 1)[1]
            ws.append([row[1]])
            ws.cell(ws.max_row, 1).font = Font(bold=True)
            ws.cell(ws.max_row, 1).fill = SECTION_FILL
            continue
        key, question, hint, choices = row
        ws.append([section, question, None, hint, key])
        ws.cell(ws.max_row, 3).fill = ANSWER_FILL
        if choices:
            drops.add(ws, ws.cell(ws.max_row, 3).coordinate, choices)
    for col, width in zip("ABCD", (22, 52, 40, 54)):
        ws.column_dimensions[col].width = width
    ws.column_dimensions["E"].hidden = True
    ws.freeze_panes = "A2"
    for r in ws.iter_rows(min_row=2):
        for c in r[:4]:
            c.alignment = WRAP

    # Row 1 (hidden): each column's key. Row 2: what to fill in. Row 3: headings. Then one row per record.
    for title, record, count, ask, fields in tables(catalogue):
        t = wb.create_sheet(title)
        t.append(["#"] + [f for f, *_ in fields])
        t.row_dimensions[1].hidden = True
        t.append([ask])
        t.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(fields) + 1)
        t.cell(2, 1).alignment = WRAP
        t.cell(2, 1).font = Font(italic=True)
        t.row_dimensions[2].height = 36
        t.append(["#"] + [h for _, h, *_ in fields])
        head(t[3])
        t.row_dimensions[3].height = 32
        for n in range(1, count + 1):
            t.append([f"{title[:-1]} {n}"])
            for c in t[t.max_row][1:len(fields) + 1]:
                c.fill = ANSWER_FILL
        t.column_dimensions["A"].width = 10
        for i, (_, _, width, choices) in enumerate(fields, 2):
            letter = t.cell(1, i).column_letter
            t.column_dimensions[letter].width = width
            if choices:
                drops.add(t, f"{letter}4:{letter}{3 + count}", choices)
        t.freeze_panes = "B4"

    fw = wb.create_sheet(FRAMEWORK_SHEET)
    fw.append(["Write this", "Framework"])
    head(fw[1])
    for f in catalogue["frameworks"]:
        fw.append([short_name(f["framework_name"]), f["framework_name"]])
    fw.column_dimensions["A"].width = 32
    fw.column_dimensions["B"].width = 70
    wb.move_sheet("Lists", offset=len(wb.sheetnames))  # keep the hidden sheet last
    wb.save(path)


MONTHS = {m: i for i, m in enumerate(["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"], 1)}


def as_date(v):
    """YYYY-MM-DD from an Excel date or text like 01-Apr-2026, 01/04/2026 or 2026-04-01; None if not a date."""
    if isinstance(v, datetime.datetime):
        return v.date().isoformat()
    if isinstance(v, datetime.date):
        return v.isoformat()
    s = str(v).strip()
    m = re.fullmatch(r"(\d{4})-(\d{1,2})-(\d{1,2})", s)
    if m:
        y, mo, d = map(int, m.groups())
    else:
        m = re.fullmatch(r"(\d{1,2})[-/ .]([A-Za-z]{3,9}|\d{1,2})[-/ .](\d{4})", s)
        if not m:
            return None
        d, mo, y = m.groups()
        mo = MONTHS.get(mo[:3].lower()) if not mo.isdigit() else int(mo)
        if not mo:
            return None
        d, y = int(d), int(y)
    try:
        return datetime.date(y, mo, d).isoformat()
    except ValueError:
        return None


def clean(v):
    if isinstance(v, str):
        return v.strip() or None
    if isinstance(v, float) and v.is_integer():
        return int(v)
    if isinstance(v, (datetime.date, datetime.datetime)):
        return as_date(v)
    return v


# ─── The client form (QA/Final_Onboarding_Questionnaire_FrameworkLogic.xlsx) ─────────────────────────────
# One sheet, "Client Onboarding Form": Category | Information Required | Client Response | Instructions. Rows
# are recognised by their label (column B). Some answers hold several details in one cell (a complete
# address, the "Additional ... (If any)" rows, scanner tool URLs); those are split up as well as text allows,
# and what cannot be placed is reported. Offices are named the way the client writes them ("HQ", "Mumbai");
# the UI matches those names to the offices.

FORM_FIELDS = {
    "company legal name": "org.customer_name",
    "category": "org.category",
    "industry": "org.industry",
    "primary web domain": "org.primary_domain",
    "office type": "office.1.office_type",
    "contract start date": "contract.contract_start_date",
    "contract period (months)": "contract.contract_period_months",
    "billing currency": "contract.currency",
    "invoicing frequency": "contract.invoicing_frequency",
    "total platform users": "licence.max_users",
    "total vcisos": "licence.max_vcisos",
    "total organisation owners": "licence.max_org_owners",
    "total it admins": "licence.max_it_admins",
    "total scanner tools": "licence.max_api_integrations",
    "total framework licenses": "licence.max_frameworks",
    "vulnerability management": "module.vulnerabilities",
    "alerting": "module.alerting",
    "mobile app access": "module.mobile_access",
    "news & breach watch": "module.news_breach_watch",
    "on-prem ai assessment": "module.onprem_ai_assessment",
    "incidents": "module.incidents",
    "login method": "security.sso_type",
    "sso metadata url": "security.sso_metadata_url",
    "session timeout (minutes)": "security.session_timeout_minutes",
    "enforce mfa for all users?": "security.mfa_enforced",
}
# "<who> - <field>" rows of the Contacts and Initial Users sections.
CONTACT_FIELDS = {"purpose": "contact_type", "priority": "contact_priority", "office location": "office",
                  "office": "office", "full name": "contact_person", "name": "contact_person",
                  "designation": "designation", "department": "department", "email": "contact_email",
                  "phone": "contact_phone", "mobile": "contact_phone"}
USER_FIELDS = {"full name": "user_name", "name": "user_name", "login id": "login_id", "email": "user_email",
               "mobile": "user_mobile", "phone": "user_mobile", "designation": "designation",
               "department": "department", "role": "role_type", "office": "office", "office location": "office"}
FRAMEWORK_FIELDS = {"applied offices": "offices", "offices": "offices",
                    "review frequency (months)": "review_months", "review frequency": "review_months"}
OFFICE_TYPES = ["Primary", "Secondary", "Data Center", "Regional Office"]
SSO = [("saml", "SAML_2_0"), ("oidc", "OIDC"), ("openid", "OIDC"), ("azure", "AZURE_AD"), ("entra", "AZURE_AD"),
       ("okta", "OKTA"), ("local", "LOCAL")]

EMAIL_RE = re.compile(r"[^@\s,;]+@[^@\s,;]+\.[^@\s,;]+")
PHONE_RE = re.compile(r"^\+?[\d\s\-()]{8,}$")
POSTCODE_RE = re.compile(r"(?<![\w-])(\d{6}|\d{5}(?:-\d{4})?|\d{4})(?![\w-])")
URL_RE = re.compile(r"https?://[^\s,;]+", re.I)


def norm(s):
    return re.sub(r"\s+", " ", str(s or "")).strip().lower()


def parts(text, seps=r"[,;\n]"):
    return [p.strip() for p in re.split(seps, str(text or "")) if p.strip()]


def office_type_in(text):
    """An office type named at the start ("Secondary: ...", "Data Center - ...") or in brackets; -> (type, rest)."""
    for t in OFFICE_TYPES:
        m = re.search(rf"\(\s*{t}( office)?\s*\)", text, re.I)
        if m:
            return t, (text[:m.start()] + text[m.end():]).strip(" ,")
        m = re.match(rf"\s*{t}( office)?\s*[:\-–]\s*", text, re.I)
        if m:
            return t, text[m.end():].strip()
    return None, text


def parse_address(text):
    """A complete address -> {office_address, city, state, country, pincode}, read from the end:
    "..., Mumbai, Maharashtra, India, 400051" (the pincode may sit anywhere, e.g. "Maharashtra 400051")."""
    out = {}
    tokens = parts(text)
    for i in range(len(tokens) - 1, -1, -1):
        m = POSTCODE_RE.search(tokens[i])
        if m:
            out["pincode"] = m.group(1)
            tokens[i] = (tokens[i][:m.start()] + tokens[i][m.end():]).strip(" -–")
            break
    else:
        # A postcode with letters (NW1 6XE, K1A 0B1), written as the last part.
        if tokens and re.fullmatch(r"(?=.*\d)(?=.*[A-Za-z])[A-Za-z0-9]{2,5}( ?[A-Za-z0-9]{2,4})?", tokens[-1]):
            out["pincode"] = tokens.pop().upper()
    tokens = [t for t in tokens if t]
    for key in ("country", "state", "city"):
        if len(tokens) > 1 or (tokens and key == "city"):
            out[key] = tokens.pop()
    if tokens:
        out["office_address"] = ", ".join(tokens)
    return out


def looks_like_office(token, offices):
    t = norm(token)
    if re.search(r"\b(office|hq|head|headquarters?|branch|primary|secondary|data cent(er|re))\b", t):
        return True
    return any(o.get("city") and norm(o["city"]) in t for o in offices.values())


def parse_person(line, opt, offices, kind):
    """One free-text person, e.g. "Vikas Singh, SOC Lead, Security Operations, vikas@x.in, +91 98..., Emergency,
    Priority 1, Bengaluru": email, phone, type/role, priority, login ID and office are recognised by their
    shape; the other parts are taken as name, designation, department in that order."""
    rec, rest = {}, []
    types = {norm(t): t for t in opt.get("contact_types" if kind == "contact" else "role_types", [])}
    designations = {norm(t) for t in opt.get("designations", [])}
    departments = {norm(t) for t in opt.get("departments", [])}
    for tok in parts(line, r"[,;|\n]"):
        n = norm(tok)
        if EMAIL_RE.fullmatch(tok) and "email" not in rec:
            rec["email"] = tok
        elif PHONE_RE.match(tok) and "phone" not in rec:
            rec["phone"] = tok
        elif n in types and "type" not in rec:
            rec["type"] = types[n]
        elif kind == "contact" and re.fullmatch(r"(priority|p)?\s*:?\s*\d", n) and "priority" not in rec:
            rec["priority"] = int(re.search(r"\d", n).group())
        elif kind == "user" and re.fullmatch(r"[a-z0-9_-]+[._][a-z0-9._-]+", n) and "login" not in rec:
            rec["login"] = tok
        elif n in departments and "department" not in rec:
            rec["department"] = tok
        elif n in designations and "designation" not in rec:
            rec["designation"] = tok
        elif looks_like_office(tok, offices) and "office" not in rec:
            rec["office"] = tok
        else:
            rest.append(tok)
    names = [k for k in ("name", "designation", "department") if k not in rec]
    for k, v in zip(names, rest):
        rec[k] = v
    keys = {"contact": {"type": "contact_type", "priority": "contact_priority", "name": "contact_person",
                        "email": "contact_email", "phone": "contact_phone"},
            "user": {"type": "role_type", "login": "login_id", "name": "user_name", "email": "user_email",
                     "phone": "user_mobile"}}[kind]
    out = {keys.get(k, k): v for k, v in rec.items()}
    return out, rest[len(names):]


def parse_framework(line):
    """One free-text framework: "DPDPA 2023 - Mumbai, Bengaluru - 6 months", "GDPR for HQ (review 12)"."""
    rec = {}
    m = re.search(r"(?:review\D{0,12})?(\d{1,2})\s*(?:months?|mo)\b|review\D{0,12}(\d{1,2})", line, re.I)
    if m:
        rec["review_months"] = m.group(1) or m.group(2)
        line = (line[:m.start()] + line[m.end():])
    line = re.sub(r"\(\s*(every)?\s*\)|\(\s*\)", "", line).strip(" ,-–()")
    # Name, then its offices after "for" / " - " / ": " (a colon inside a name, as in 27001:2022, has no space after).
    m = re.split(r"\s+(?:for|at|in|on)\s+|\s+[-–]\s+|\s*:\s+", line, maxsplit=1, flags=re.I)
    if len(m) == 2:
        rec["name"], rec["offices"] = m[0].strip(), m[1].strip(" -–()")
    else:
        rec["name"] = line
    return rec


def read_form(ws, opt):
    answers, warnings, notes = {}, [], []
    offices = {}                       # number -> parsed office
    groups = {"contact": {}, "user": {}, "framework": {}}  # prefix -> record, in order of appearance
    extra = []                         # (kind, text) of the "Additional ... (If any)" rows
    section = ""
    for row in ws.iter_rows(min_row=2, values_only=True):
        row = list(row) + [None] * 4
        category, label, value = row[0], row[1], clean(row[2])
        if category and not label:
            section = norm(re.sub(r"^\d+\.\s*", "", str(category)))
            continue
        if not label:
            continue
        section = norm(category) or section
        lab = norm(label)
        if value is None:
            continue
        if lab.startswith("additional"):
            kind = next((k for k in ("office", "contact", "framework", "user") if k in lab), None)
            if kind:
                extra.append((kind, str(value)))
            continue
        if lab == "complete address":
            otype, text = office_type_in(str(value))
            offices.setdefault(1, {}).update(parse_address(text))
            if otype:
                offices[1].setdefault("office_type", otype)
            continue
        if lab == "office type":
            offices.setdefault(1, {})["office_type"] = str(value)
            continue
        if lab == "names & domains to monitor":
            toks = parts(value)
            dom = [t for t in toks if re.fullmatch(r"[\w-]+(\.[\w-]+)+", t)]
            answers["news.monitored_domains"] = ", ".join(dom)
            answers["news.monitored_names"] = ", ".join(t for t in toks if t not in dom)
            continue
        if lab == "scanner tool urls":
            text = str(value)
            for tool in opt.get("tools", []):
                m = re.search(rf"{re.escape(tool)}\W{{0,5}}(https?://[^\s,;]+)?", text, re.I)
                url = m.group(1) if m else next((u for u in URL_RE.findall(text) if tool.lower() in u.lower()), None)
                if m or url:
                    answers[f"tool.{tool}.used"] = "Yes"
                    if url:
                        answers[f"tool.{tool}.endpoint"] = url
            continue
        if lab == "login method":
            answers["security.sso_type"] = next((v for k, v in SSO if k in norm(value)), str(value))
            continue
        if lab in FORM_FIELDS:
            answers[FORM_FIELDS[lab]] = value
            continue
        # "<who> - <field>" in Contacts / Initial Users; "Framework 1" and "Framework 1 - <field>".
        who, sep, field = str(label).rpartition(" - ")
        who = norm(re.sub(r"\(.*?\)", "", who or label))
        field = norm(field)
        if "framework" in section:
            rec = groups["framework"].setdefault(who, {})
            if not sep:  # "Framework 1" itself: the framework's name
                rec["name"] = value
            elif field in FRAMEWORK_FIELDS:
                rec[FRAMEWORK_FIELDS[field]] = value
            continue
        kind = "contact" if "contact" in section else "user" if "user" in section else None
        names = CONTACT_FIELDS if kind == "contact" else USER_FIELDS
        if kind and field in names:
            groups[kind].setdefault(who, {})[names[field]] = value
            continue
        notes.append(f'"{label}" is not a question this upload knows, left out')

    for kind, text in extra:
        lines = parts(text, r"[\n;]") if kind != "office" else parts(text, r"\n|;|\|")
        for line in lines:
            if kind == "office":
                otype, rest = office_type_in(line)
                rec = parse_address(rest)
                if otype:
                    rec["office_type"] = otype
                offices[max(offices, default=0) + 1] = rec
            elif kind == "framework":
                groups["framework"][f"extra {len(groups['framework'])}"] = parse_framework(line)
            else:
                rec, left = parse_person(line, opt, offices, kind)
                groups[kind][f"extra {len(groups[kind])}"] = rec
                if left:
                    notes.append(f'Additional {kind}s: "{", ".join(left)}" could not be placed')

    for n, rec in offices.items():
        for k, v in rec.items():
            answers[f"office.{n}.{k}"] = v
    for kind, recs in groups.items():
        for n, rec in enumerate(recs.values(), 1):
            for k, v in rec.items():
                answers[f"{kind}.{n}.{k}"] = v
    return answers, warnings + notes


def read_answers(path, opt=None):
    wb = load_workbook(path, data_only=True)
    if QA_SHEET in wb.sheetnames:
        ws = wb[QA_SHEET]
    elif OLD_SHEET in wb.sheetnames and wb[OLD_SHEET].cell(1, 5).value == "Key":
        ws = wb[OLD_SHEET]  # the earlier single-sheet template: same keys, offices numbered in rows
    elif OLD_SHEET in wb.sheetnames:
        answers, warnings = read_form(wb[OLD_SHEET], opt or {})
        out = {}
        for k, v in answers.items():
            if k.endswith("_date") and not re.fullmatch(r"\d{4}-\d{2}-\d{2}", str(v)):
                d = as_date(v)
                if d is None:
                    warnings.append(f'Contract start date: "{v}" is not a date (use DD-MMM-YYYY), left empty')
                    continue
                v = d
            out[k] = v
        if not out:
            warnings.append("The questionnaire has no answers filled in")
        return {"answers": out, "warnings": warnings}
    else:
        raise ValueError(f'This is not the onboarding questionnaire: it has no "{QA_SHEET}" or "{OLD_SHEET}" sheet.')
    answers, warnings = {}, []

    def put(key, label, value):
        value = clean(value)
        if value is None:
            return
        if key.endswith("_date") and not re.fullmatch(r"\d{4}-\d{2}-\d{2}", str(value)):
            d = as_date(value)
            if d is None:
                warnings.append(f'{label}: "{value}" is not a date (use DD-MMM-YYYY), left empty')
                return
            value = d
        answers[key] = value

    for row in ws.iter_rows(min_row=2, values_only=True):
        row = list(row) + [None] * 5
        if row[4]:
            put(row[4], row[1], row[2])

    # The tables: keys in the hidden first row, records from row 4 down. A record keeps its row's number
    # ("Office 2" is the second row however many are blank), which is how Contacts and Users name offices.
    for title, record in TABLES:
        if title not in wb.sheetnames:
            continue
        t = wb[title]
        keys = [c.value for c in t[1]]
        for n, row in enumerate(t.iter_rows(min_row=4, values_only=True), 1):
            for k, v in zip(keys[1:], row[1:]):
                if k:
                    put(f"{record}.{n}.{k}", f"{title}, {title[:-1]} {n}, {k}", v)

    if not answers:
        warnings.append("The questionnaire has no answers filled in")
    return {"answers": answers, "warnings": warnings}


if __name__ == "__main__":
    # Used by ctw-admin (lib/python.js): `template <out.xlsx>` with the catalogue JSON on stdin, or
    # `read` with {"file": <base64 .xlsx>, "options": <the choice lists>} on stdin. Prints JSON.
    import os
    import tempfile
    cmd = sys.argv[1] if len(sys.argv) > 1 else ""
    data = json.load(sys.stdin)
    if cmd == "template":
        write_template(sys.argv[2], data)
        print(json.dumps({"ok": True}))
    elif cmd == "read":
        fd, tmp = tempfile.mkstemp(suffix=".xlsx")
        try:
            with os.fdopen(fd, "wb") as f:
                f.write(base64.b64decode(data["file"]))
            try:
                print(json.dumps(read_answers(tmp, data.get("options"))))
            except ValueError as e:
                print(json.dumps({"error": str(e)}))
            except Exception:
                print(json.dumps({"error": "The file could not be read as an Excel (.xlsx) questionnaire"}))
        finally:
            os.remove(tmp)
    else:
        sys.exit("usage: questionnaire.py template <out.xlsx> | read  (JSON on stdin)")
